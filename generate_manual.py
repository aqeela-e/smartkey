import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    return heading

def add_screenshot_placeholder(doc, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"\n[ --- MINTA TOLONG TEMPELKAN SCREENSHOT: {caption_text} DI SINI --- ]\n")
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.bold = True
    
    caption = doc.add_paragraph(f"Gambar: {caption_text}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.style = 'Caption'

def main():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('Buku Panduan Penggunaan Aplikasi SmartKey', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Analisis Optimasi Parameter AI pada Digital Signature')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # Section 1
    add_heading(doc, '1. Pengenalan Aplikasi', level=1)
    doc.add_paragraph(
        "Aplikasi SmartKey adalah purwarupa (Proof-of-Concept) yang dibangun menggunakan kerangka kerja antarmuka "
        "Streamlit. Aplikasi ini digunakan untuk membangkitkan kunci kriptografi berdasarkan algoritma Kecerdasan "
        "Buatan (Artificial Intelligence) seperti Genetic Algorithm (GA), Artificial Neural Network (ANN), dan "
        "Ant Colony Optimization (ACO). Aplikasi ini juga dilengkapi dengan alat uji statistik NIST dan kalkulator Entropi."
    )
    
    # Section 2
    add_heading(doc, '2. Akses Aplikasi', level=1)
    doc.add_paragraph("Langkah-langkah untuk menjalankan aplikasi:")
    doc.add_paragraph("1. Pastikan Anda telah mengaktifkan Virtual Environment.", style='List Number')
    doc.add_paragraph("2. Jalankan perintah `streamlit run app.py` di terminal.", style='List Number')
    doc.add_paragraph("3. Aplikasi akan otomatis terbuka di peramban web pada alamat http://localhost:8501.", style='List Number')
    
    add_screenshot_placeholder(doc, "Tampilan Awal Aplikasi SmartKey di Web Browser")
    
    # Section 3
    add_heading(doc, '3. Panel Konfigurasi (Sidebar)', level=1)
    doc.add_paragraph(
        "Di sebelah kiri aplikasi, terdapat Sidebar yang berfungsi sebagai pusat kendali untuk mengatur parameter "
        "Algoritma AI dan Kriptografi."
    )
    
    add_heading(doc, '3.1 Konfigurasi AI', level=2)
    doc.add_paragraph(
        "Anda dapat memilih salah satu dari tiga algoritma AI (GA, ANN, atau ACO). Setiap kali Anda mengubah pilihan, "
        "panel parameter di bawahnya akan menyesuaikan diri secara otomatis."
    )
    doc.add_paragraph("• Genetic Algorithm (GA): Anda dapat mengubah Population Size, Generations, Mutation Rate, dll.", style='List Bullet')
    doc.add_paragraph("• Artificial Neural Network (ANN): Mengatur jumlah node pada Hidden Layer dan deviasi noise.", style='List Bullet')
    doc.add_paragraph("• Ant Colony Optimization (ACO): Mengatur jumlah semut, iterasi, dan rasio evaporasi.", style='List Bullet')
    
    add_screenshot_placeholder(doc, "Sidebar - Pemilihan Algoritma AI dan Parameternya")
    
    add_heading(doc, '3.2 Konfigurasi Kriptografi', level=2)
    doc.add_paragraph(
        "Pilih jenis algoritma (RSA, DSA, ElGamal) dan ukuran bits kunci (contoh: 2048) yang akan digunakan untuk proses Tanda Tangan Digital."
    )
    
    add_screenshot_placeholder(doc, "Sidebar - Pemilihan Algoritma Kriptografi")
    
    # Section 4
    add_heading(doc, '4. Proses Penandatanganan (Signing)', level=1)
    doc.add_paragraph(
        "Proses utama dilakukan di panel tengah layar. Terdapat dua mode penandatanganan yang bisa Anda pilih melalui tab, yaitu Manual Mode dan Auto-Optimizer Mode."
    )
    
    add_heading(doc, '4.1 Mode Manual', level=2)
    doc.add_paragraph("1. Klik tombol 'Browse Files' atau seret dokumen ke area unggah.", style='List Number')
    doc.add_paragraph("2. Aplikasi akan otomatis menghitung nilai Hash (SHA-256) dari dokumen.", style='List Number')
    doc.add_paragraph("3. Pastikan Anda telah memilih AI di Sidebar, lalu klik tombol biru 'Tandatangani Secara Manual'.", style='List Number')
    doc.add_paragraph("4. Algoritma AI akan dijalankan dan kuncinya digunakan untuk penandatanganan.", style='List Number')
    
    add_heading(doc, '4.2 Mode Auto-Optimizer', level=2)
    doc.add_paragraph(
        "Mode ini digunakan untuk mendapatkan konfigurasi AI paling optimal (Entropi Tertinggi). "
        "Langkah-langkahnya:"
    )
    doc.add_paragraph("1. Setelah mengunggah dokumen, pindah ke tab 'Auto-Optimizer Mode'.", style='List Number')
    doc.add_paragraph("2. Klik tombol 'Cari Konfigurasi Terbaik & Sign'.", style='List Number')
    doc.add_paragraph(
        "3. Sistem akan secara otomatis menjalankan ketiga algoritma AI (GA, ANN, dan ACO) di latar belakang.", style='List Number'
    )
    doc.add_paragraph(
        "4. Sistem memilih AI dengan skor tertinggi untuk kemudian dikawinkan dengan algoritma Kriptografi yang Anda pilih di Sidebar untuk ditandatangani.", style='List Number'
    )
    
    add_screenshot_placeholder(doc, "Proses Tanda Tangan menggunakan Auto-Optimizer Mode")
    
    # Section 5
    add_heading(doc, '5. Analisis Dashboard & Evaluasi Kunci', level=1)
    doc.add_paragraph(
        "Setelah penandatanganan selesai, gulir layar ke bawah untuk melihat hasil evaluasi keacakan (Entropi) "
        "dari seed AI yang baru saja di-generate."
    )
    
    doc.add_paragraph("Dashboard menampilkan 4 metrik utama:")
    doc.add_paragraph("• Waktu Generate AI (kecepatan).", style='List Bullet')
    doc.add_paragraph("• Entropi Shannon (makin tinggi makin baik).", style='List Bullet')
    doc.add_paragraph("• Normalisasi Entropi (skala 0 - 1).", style='List Bullet')
    doc.add_paragraph("• NIST Pass Rate (persentase kelulusan uji acak).", style='List Bullet')
    
    add_screenshot_placeholder(doc, "Tabel Metrik Utama dan Hasil Uji NIST")
    
    doc.add_paragraph(
        "Terdapat pula visualisasi berupa Pie Chart yang merangkum kelulusan uji NIST, serta Progress Bar untuk 'Skor Kualitas Kunci'."
    )
    
    add_heading(doc, '5.1 Leaderboard Algoritma', level=2)
    doc.add_paragraph(
        "Jika Anda menggunakan Auto-Optimizer Mode, akan muncul tabel 'Leaderboard Algoritma' tambahan di bagian bawah "
        "dashboard. Tabel ini memperlihatkan perbandingan langsung antara GA, ANN, dan ACO. Algoritma dengan Combined Score "
        "tertinggi akan mendapat predikat 'WINNER' dan otomatis digunakan oleh sistem."
    )
    add_screenshot_placeholder(doc, "Leaderboard Algoritma (Hasil Perbandingan Auto-Optimizer)")
    
    # Section 6
    add_heading(doc, '6. Proses Verifikasi', level=1)
    doc.add_paragraph(
        "Untuk memastikan Tanda Tangan valid:"
    )
    doc.add_paragraph("1. Pastikan Anda masih berada di sesi yang sama.", style='List Number')
    doc.add_paragraph("2. Klik tombol 'Verifikasi Tanda Tangan (Verify)'.", style='List Number')
    doc.add_paragraph("3. Jika dokumen belum diubah, Anda akan melihat pesan hijau yang menyatakan dokumen Valid dan Asli.", style='List Number')
    
    add_screenshot_placeholder(doc, "Pesan Sukses Verifikasi Dokumen")
    
    # Section 7
    add_heading(doc, '7. Ekspor Laporan', level=1)
    doc.add_paragraph(
        "Di bagian paling bawah, Anda dapat mengunduh seluruh rangkuman eksperimen."
    )
    doc.add_paragraph("• Unduh Laporan JSON: Berisi metrik waktu, entropi, parameter, dan hasil uji NIST secara lengkap.", style='List Bullet')
    doc.add_paragraph("• Unduh File Signature (.sig): Berkas mentah dari tanda tangan kriptografinya.", style='List Bullet')
    
    add_screenshot_placeholder(doc, "Panel Unduh Laporan dan Signature")
    
    doc.save('d:/SmartKey/Panduan_Penggunaan_SmartKey.docx')
    print("Document saved.")

if __name__ == '__main__':
    main()
